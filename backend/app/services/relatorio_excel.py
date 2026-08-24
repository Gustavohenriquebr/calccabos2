from collections import defaultdict
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.calculo import CONTEXTOS
from app.services.aterramento import calcular_aterramento
from app.services.areas_classificadas import calcular_areas_classificadas
from app.services.para_raios import calcular_para_raios
from app.services.projeto_eletrico import calcular_sistema_trifasico_projeto, calcular_transformador, carregar_json
from app.services.protecao import validar_protecao_geral
from app.services.status_utils import limpar_texto, normalizar_status, normalizar_tipo_cabo, status_global, status_visual
from app.services.memorial_cabos import montar_memorial_cabos


# ─── Paleta de cores ──────────────────────────────────────────────────────────
PETROBRAS_BLUE  = "003087"
ACCENT_BLUE     = "1D4ED8"
LIGHT_BLUE      = "EAF1FB"
LIGHT_GRAY      = "F5F7FA"
WHITE           = "FFFFFF"
OK_FILL         = "D1FAE5"
OK_TEXT         = "065F46"
ALERT_FILL      = "FEF3C7"
ALERT_TEXT      = "92400E"
CRITICAL_FILL   = "FEE2E2"
CRITICAL_TEXT   = "991B1B"
GRID            = "CBD5E1"
DARK_NAVY       = "0F172A"
SLATE           = "334155"
SURFACE         = "F8FAFC"
MEDIUM_GRAY     = "E2E8F0"
HEADER_ROW      = "1E3A5F"   # azul escuro profissional para cabeçalhos de tabela
SUBHEADER_ROW   = "2D5282"   # azul médio para sub-cabeçalhos


HEADERS = [
    "CIRCUIT", "FROM", "TO", "PROT.DEVICE", "Vn", "PHASE", "AC/DC",
    "kW", "kVA", "PF", "η", "n", "L(m)", "Type of cable", "install.",
    "K1", "K2", "K3", "Ib", "Ib'", "ICOND", "CP×Ø", "CK",
    "RDC", "RAC", "XL", "ΔVmax", "ΔVcalc", "CK", "Isc_LOC",
    "Isc_CAB", "t(s)", "FORMATION", "FINAL CONTROL", "CABLE TYPE",
    "LENGTH", "REV", "NOTE",
]

# Índices (1-based) das colunas críticas a destacar na aba Memória de Cálculo
CRITICAL_COLS = {
    19,  # Ib
    20,  # Ib'
    21,  # ICOND
    30,  # Isc_LOC
    28,  # ΔVcalc
    34,  # FINAL CONTROL
    35,  # CABLE TYPE
}


# ─── Helpers básicos ──────────────────────────────────────────────────────────

def _ctx(contexto):
    if hasattr(contexto, "value"):
        return contexto.value
    return str(contexto or "industrial").split(".")[-1]


def _tipo_cabo(tipo):
    return normalizar_tipo_cabo(tipo)


def _status_final(c):
    if getattr(c, "status_final", None):
        return normalizar_status(c.status_final)
    if getattr(c, "status", None) == "ok":
        return "OK"
    if getattr(c, "status", None) == "erro":
        return "CRITICO"
    if getattr(c, "status", None) == "alerta":
        return "ALERTA"
    return "ALERTA"


def _display_status(status):
    """Garante exibição de CRÍTICO com acento."""
    st = normalizar_status(status, "ALERTA")
    if st == "CRITICO":
        return "CRÍTICO"
    return st


def _status_fill(status):
    status = normalizar_status(status)
    if status == "OK":
        return OK_FILL
    if status == "CRITICO":
        return CRITICAL_FILL
    return ALERT_FILL


def _status_text_color(status):
    status = normalizar_status(status)
    if status == "OK":
        return OK_TEXT
    if status == "CRITICO":
        return CRITICAL_TEXT
    return ALERT_TEXT


def _contagem_protecoes_circuitos(circuitos):
    statuses = [normalizar_status(getattr(c, "protecao_status", None), "OK") for c in circuitos]
    return {
        "ok": sum(1 for status in statuses if status == "OK"),
        "alerta": sum(1 for status in statuses if status == "ALERTA"),
        "critico": sum(1 for status in statuses if status == "CRITICO"),
    }


def _resumo_global(projeto, circuitos, modulos=None):
    circuito_status = [_status_final(c) for c in circuitos]
    if modulos is None:
        transformador = calcular_transformador(carregar_json(getattr(projeto, "transformador_dados", None)))
        modulos = {
            "Proteção geral": validar_protecao_geral(carregar_json(getattr(projeto, "protecao_geral_dados", None)), transformador),
            "Para-raios": calcular_para_raios(carregar_json(getattr(projeto, "para_raios_dados", None))),
            "Aterramento": calcular_aterramento(carregar_json(getattr(projeto, "aterramento_dados", None))),
            "Áreas Classificadas": calcular_areas_classificadas(carregar_json(getattr(projeto, "areas_classificadas_dados", None))),
        }
    status_items = circuito_status
    status_items.extend(normalizar_status(getattr(c, "protecao_status", None), "OK") for c in circuitos)
    status_items.extend(normalizar_status(dados.get("status"), "ALERTA") for dados in modulos.values())
    geral = status_global(*status_items)
    criticos = sum(1 for status in status_items if normalizar_status(status, "ALERTA") == "CRITICO")
    alertas = sum(1 for status in status_items if normalizar_status(status, "ALERTA") == "ALERTA")
    return {
        "status": geral,
        "protecoes": _contagem_protecoes_circuitos(circuitos),
        "modulos": modulos,
        "modulos_contagem": {
            "ok": sum(1 for dados in modulos.values() if normalizar_status(dados.get("status"), "ALERTA") == "OK"),
            "alerta": sum(1 for dados in modulos.values() if normalizar_status(dados.get("status"), "ALERTA") == "ALERTA"),
            "critico": sum(1 for dados in modulos.values() if normalizar_status(dados.get("status"), "ALERTA") == "CRITICO"),
        },
        "motivo_status": (
            f"Há {criticos} pendência(s) crítica(s) e {alertas} pendência(s) em alerta."
            if criticos
            else f"Há {alertas} pendência(s) em alerta."
            if alertas
            else "Todos os módulos avaliados estão OK."
        ),
    }


def _excel_value(valor):
    if valor is None or valor == "":
        return "N/D"
    bruto = str(valor).strip().upper().replace("_", " ")
    if bruto in {"OK", "ALERTA", "CRITICO", "CRÍTICO", "NAO CALCULADO", "NÃO CALCULADO"}:
        return status_visual(valor)
    if isinstance(valor, str):
        return limpar_texto(valor)
    return valor


# ─── Estilos básicos ──────────────────────────────────────────────────────────

def _thin_border():
    side = Side(style="thin", color=GRID)
    return Border(left=side, right=side, top=side, bottom=side)


def _medium_border():
    side = Side(style="medium", color=SLATE)
    return Border(left=side, right=side, top=side, bottom=side)


def _bottom_border_only():
    bottom = Side(style="thin", color=GRID)
    return Border(bottom=bottom)


def _header(cell, bg=HEADER_ROW, font_size=9, bold=True):
    """Cabeçalho de tabela — fundo escuro, texto branco."""
    cell.font = Font(bold=bold, color=WHITE, size=font_size, name="Calibri")
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = _thin_border()


def _body(cell, fill=None, bold=False, align="center", color=SLATE, size=9):
    """Célula de corpo — fundo opcional, texto escuro."""
    cell.font = Font(size=size, color=color, bold=bold, name="Calibri")
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = _thin_border()
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)


def _status_cell(cell, status, bold=True, size=9):
    """Aplica cor de fundo e texto de acordo com o status."""
    fill = _status_fill(status)
    text_color = _status_text_color(status)
    cell.font = Font(bold=bold, color=text_color, size=size, name="Calibri")
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = _thin_border()


def _section_header(ws, row, col_start, col_end, titulo, bg=DARK_NAVY):
    """Faixa de título de seção em fundo escuro."""
    ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)
    cell = ws.cell(row=row, column=col_start, value=titulo)
    cell.font = Font(bold=True, color=WHITE, size=10, name="Calibri")
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 20


def _ajustar_larguras(ws, larguras):
    for index, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(index)].width = largura


def _titulo_aba(ws, titulo, subtitulo, ultima_coluna):
    """Cabeçalho de aba com título e subtítulo."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ultima_coluna)
    ws.cell(row=1, column=1, value=titulo)
    ws.cell(row=1, column=1).font = Font(bold=True, color=WHITE, size=14, name="Calibri")
    ws.cell(row=1, column=1).fill = PatternFill("solid", fgColor=DARK_NAVY)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    if subtitulo:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ultima_coluna)
        ws.cell(row=2, column=1, value=subtitulo)
        ws.cell(row=2, column=1).font = Font(color=SLATE, size=9, name="Calibri")
        ws.cell(row=2, column=1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(row=2, column=1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[2].height = 18


# ─── Linha de dados do circuito ───────────────────────────────────────────────

def _linha(c):
    qt_calc = c.queda_tensao_acumulada if c.queda_tensao_acumulada is not None else c.queda_tensao_pct
    return [
        c.tag or c.descricao,
        c.from_barramento or "",
        c.to_equipamento or "",
        c.protection_device or "",
        c.tensao,
        c.fases,
        c.corrente_ac_dc or "AC",
        c.potencia_kw,
        c.potencia_kva,
        c.fator_potencia,
        c.fator_eficiencia,
        c.fator_demanda,
        c.distancia_m,
        _tipo_cabo(c.tipo_cabo),
        c.metodo_instalacao or "",
        c.fator_k1,
        c.fator_k2,
        c.fator_k3,
        c.corrente_projeto or c.corrente_nominal,
        None,
        c.corrente_condutor or c.ampacidade,
        c.secao_mm2,
        None,
        c.impedancia_rdc,
        c.impedancia_rac,
        c.impedancia_xl,
        c.queda_tensao_max,
        qt_calc,
        None,
        c.isc_local,
        c.isc_cabo,
        c.tempo_atuacao,
        c.formacao,
        status_visual(_status_final(c)),
        c.tipo_cabo_comercial,
        c.comprimento_real or c.distancia_m,
        c.revisao or "0",
        c.nota_tecnica or "",
    ]


# ─── Aba Memória de Cálculo ───────────────────────────────────────────────────

def _aba_memoria(wb, projeto, circuitos, usuario):
    ws = wb.active
    ws.title = "Memória de Cálculo"
    ws.sheet_properties.tabColor = DARK_NAVY
    ws.freeze_panes = "A9"

    # ── Título e subtítulo ────────────────────────────────────────────────────
    num_cols = len(HEADERS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    ws["A1"] = "CALCCABOS  ·  MEMORY CALCULATION FOR LV/MV CABLES"
    ws["A1"].font = Font(bold=True, color=WHITE, size=14, name="Calibri")
    ws["A1"].fill = PatternFill("solid", fgColor=DARK_NAVY)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30

    # ── Bloco de metadados ────────────────────────────────────────────────────
    contexto = CONTEXTOS.get(_ctx(projeto.contexto), CONTEXTOS["industrial"])
    meta_esquerda = [
        ("Projeto", projeto.nome),
        ("Cliente", projeto.cliente or "-"),
        ("Engenheiro responsável", usuario.nome),
        ("CREA", usuario.crea or "-"),
    ]
    meta_direita = [
        ("Empresa", usuario.empresa or "-"),
        ("Contexto", _ctx(projeto.contexto)),
        ("Normas", f"NBR 5410 / IEC 60909 / {contexto['norma_ref']}"),
        ("Data", datetime.now().strftime("%d/%m/%Y %H:%M")),
    ]
    for idx, (label, value) in enumerate(meta_esquerda, start=2):
        lbl = ws.cell(row=idx, column=1, value=label)
        lbl.font = Font(bold=True, color=PETROBRAS_BLUE, size=9, name="Calibri")
        lbl.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        lbl.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        val = ws.cell(row=idx, column=2, value=_excel_value(value))
        val.font = Font(size=9, name="Calibri")
        val.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=5)

    for idx, (label, value) in enumerate(meta_direita, start=2):
        lbl = ws.cell(row=idx, column=7, value=label)
        lbl.font = Font(bold=True, color=PETROBRAS_BLUE, size=9, name="Calibri")
        lbl.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        lbl.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        val = ws.cell(row=idx, column=8, value=_excel_value(value))
        val.font = Font(size=9, name="Calibri")
        val.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=idx, start_column=8, end_row=idx, end_column=11)

    # ── Linha separadora antes do cabeçalho ───────────────────────────────────
    ws.row_dimensions[7].height = 6

    # ── Cabeçalho da tabela ───────────────────────────────────────────────────
    header_row = 8
    ws.row_dimensions[header_row].height = 36
    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=header_row, column=col, value=header)
        bg = SUBHEADER_ROW if col in CRITICAL_COLS else HEADER_ROW
        _header(cell, bg=bg, font_size=8)

    # ── Linhas de dados ───────────────────────────────────────────────────────
    for i, c in enumerate(circuitos, start=1):
        row = header_row + i
        ws.row_dimensions[row].height = 16
        status = _status_final(c)
        row_fill = LIGHT_GRAY if i % 2 == 0 else WHITE

        for col, value in enumerate(_linha(c), 1):
            cell = ws.cell(row=row, column=col, value=_excel_value(value))
            if col in CRITICAL_COLS and col != 34:
                _body(cell, fill=row_fill, bold=True, color=DARK_NAVY)
            else:
                _body(cell, fill=row_fill)

        # Fórmulas Excel
        ws.cell(row=row, column=20, value=f'=IFERROR(S{row}/(P{row}*Q{row}*R{row}),"")')
        ws.cell(row=row, column=23, value=f'=IF(U{row}>=T{row}/MAX(AG{row},1),"OK","NOK")')
        ws.cell(
            row=row,
            column=28,
            value=(
                f'=IFERROR(IF(G{row}="DC",'
                f'2*S{row}*(X{row}/MAX(AG{row},1))*(AJ{row}/1000)/E{row}*100,'
                f'IF(F{row}=3,'
                f'SQRT(3)*S{row}*((Y{row}/MAX(AG{row},1))*J{row}+(Z{row}/MAX(AG{row},1))*SQRT(MAX(0,1-J{row}^2)))*(AJ{row}/1000)/E{row}*100,'
                f'2*S{row}*((Y{row}/MAX(AG{row},1))*J{row}+(Z{row}/MAX(AG{row},1))*SQRT(MAX(0,1-J{row}^2)))*(AJ{row}/1000)/E{row}*100)), "")'
            ),
        )
        ws.cell(row=row, column=29, value=f'=IF(AB{row}<=AA{row},"OK","NOK")')
        ws.cell(row=row, column=34, value=f'=IF(OR(W{row}="NOK",AC{row}="NOK"),"CRÍTICO",IF(AD{row}="","ALERTA","OK"))')

        # Estilo das células de fórmula
        for col in (20, 23, 28, 29):
            _body(ws.cell(row=row, column=col), fill=row_fill, bold=True, color=DARK_NAVY)

        # Status final: cor semântica
        status_cell = ws.cell(row=row, column=34)
        _status_cell(status_cell, status, bold=True)

    # ── Auto-filtro na linha de cabeçalho ─────────────────────────────────────
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(num_cols)}{header_row}"

    _ajustar_larguras(
        ws,
        [16, 18, 18, 13, 9, 8, 8, 9, 9, 8, 8, 8, 9, 13, 11, 8, 8, 8, 10, 10,
         10, 9, 8, 11, 11, 10, 10, 11, 8, 10, 10, 8, 11, 14, 14, 10, 8, 24],
    )
    return ws


# ─── Aba Lista de Cargas ──────────────────────────────────────────────────────

def _aba_lista_cargas(wb, circuitos):
    ws = wb.create_sheet("Lista de Cargas")
    ws.sheet_properties.tabColor = PETROBRAS_BLUE
    _titulo_aba(ws, "Lista de Cargas", "Resumo das cargas cadastradas no projeto.", 9)

    headers = ["TAG", "Descrição", "Barramento", "Equipamento", "kW", "kVA", "FP", "η", "Demanda"]
    header_row = 4
    ws.row_dimensions[header_row].height = 26
    for col, header in enumerate(headers, 1):
        _header(ws.cell(row=header_row, column=col, value=header))

    for row_num, c in enumerate(circuitos, start=5):
        ws.row_dimensions[row_num].height = 16
        row_fill = LIGHT_GRAY if (row_num - 4) % 2 == 0 else WHITE
        values = [
            c.tag or "",
            c.descricao,
            c.from_barramento or "",
            c.to_equipamento or "",
            c.potencia_kw,
            c.potencia_kva,
            c.fator_potencia,
            c.fator_eficiencia,
            c.fator_demanda,
        ]
        for col, value in enumerate(values, 1):
            align = "left" if col <= 4 else "center"
            _body(ws.cell(row=row_num, column=col, value=_excel_value(value)), fill=row_fill, align=align)

    ws.auto_filter.ref = f"A{header_row}:I{header_row}"
    _ajustar_larguras(ws, [16, 30, 18, 22, 10, 10, 8, 8, 10])
    ws.freeze_panes = "A5"
    return ws


# ─── Aba Resumo Executivo ─────────────────────────────────────────────────────

def _aba_resumo(wb, projeto, circuitos):
    ws = wb.create_sheet("Resumo Executivo")
    ws.sheet_properties.tabColor = DARK_NAVY
    _titulo_aba(ws, "Resumo Executivo — CalcCabos", "Status consolidado, distribuição por barramento e motivo do status geral.", 8)

    resumo_global = _resumo_global(projeto, circuitos)
    status_geral = resumo_global["status"]
    status_disp = _display_status(status_geral)

    por_barramento = defaultdict(lambda: {"total": 0, "kw": 0.0, "kva": 0.0})
    por_status = defaultdict(int)
    for c in circuitos:
        barramento = c.from_barramento or "Sem barramento"
        por_barramento[barramento]["total"] += 1
        por_barramento[barramento]["kw"] += c.potencia_kw or 0
        por_barramento[barramento]["kva"] += c.potencia_kva or 0
        por_status[normalizar_status(_status_final(c), "ALERTA")] += 1

    # ── Faixa de status geral ─────────────────────────────────────────────────
    ws.merge_cells("A3:H3")
    if status_geral == "CRITICO":
        status_texto = f"⚠  STATUS GERAL: {status_disp}  —  PROJETO NÃO LIBERADO PARA EMISSÃO FINAL"
    elif status_geral == "ALERTA":
        status_texto = f"▲  STATUS GERAL: {status_disp}  —  PROJETO REQUER REVISÃO TÉCNICA"
    else:
        status_texto = f"✔  STATUS GERAL: {status_disp}  —  PROJETO APROVADO"
    ws["A3"] = status_texto
    ws["A3"].font = Font(bold=True, color=_status_text_color(status_geral), size=11, name="Calibri")
    ws["A3"].fill = PatternFill("solid", fgColor=_status_fill(status_geral))
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 28

    # ── KPIs numéricos ────────────────────────────────────────────────────────
    kpis = [
        ("Total circuitos", len(circuitos), None),
        ("OK", por_status.get("OK", 0), "OK"),
        ("Alertas", por_status.get("ALERTA", 0), "ALERTA"),
        ("Críticos", por_status.get("CRITICO", 0), "CRITICO"),
        ("Módulos OK", resumo_global["modulos_contagem"]["ok"], "OK"),
        ("Módulos alerta", resumo_global["modulos_contagem"]["alerta"], "ALERTA" if resumo_global["modulos_contagem"]["alerta"] else None),
        ("Módulos críticos", resumo_global["modulos_contagem"]["critico"], "CRITICO" if resumo_global["modulos_contagem"]["critico"] else None),
        ("Status geral", status_disp, status_geral),
    ]
    ws.row_dimensions[5].height = 20
    ws.row_dimensions[6].height = 28
    for idx, (label, value, status) in enumerate(kpis, start=1):
        lbl_cell = ws.cell(row=5, column=idx, value=label)
        _header(lbl_cell, bg=HEADER_ROW, font_size=8)
        val_cell = ws.cell(row=6, column=idx, value=_excel_value(value))
        if status:
            _status_cell(val_cell, status, bold=True, size=11)
        else:
            _body(val_cell, fill=LIGHT_BLUE, bold=True, size=11, color=DARK_NAVY)

    # ── Motivo do status geral ────────────────────────────────────────────────
    ws.row_dimensions[8].height = 18
    ws.row_dimensions[9].height = 32
    ws.merge_cells("A8:H8")
    _section_header(ws, 8, 1, 8, "Motivo do status geral")
    ws.merge_cells("A9:H9")
    ws["A9"] = resumo_global["motivo_status"]
    ws["A9"].font = Font(size=9, color=DARK_NAVY, name="Calibri")
    ws["A9"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws["A9"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)

    # ── Distribuição por barramento ───────────────────────────────────────────
    ws.row_dimensions[11].height = 18
    _section_header(ws, 11, 1, 4, "Distribuição por barramento")
    ws.row_dimensions[12].height = 22
    for col, header in enumerate(["Barramento", "Circuitos", "kW total", "kVA total"], 1):
        _header(ws.cell(row=12, column=col, value=header))
    for row_idx, (barramento, dados) in enumerate(sorted(por_barramento.items()), start=13):
        ws.row_dimensions[row_idx].height = 16
        row_fill = LIGHT_GRAY if (row_idx - 12) % 2 == 0 else WHITE
        values = [_excel_value(barramento), dados["total"], round(dados["kw"], 3), round(dados["kva"], 3)]
        for col, value in enumerate(values, 1):
            align = "left" if col == 1 else "center"
            _body(ws.cell(row=row_idx, column=col, value=value), fill=row_fill, align=align)

    ws.auto_filter.ref = "A12:D12"
    ws.freeze_panes = "A5"
    _ajustar_larguras(ws, [24, 12, 12, 12, 14, 14, 16, 18])
    return ws


# ─── Aba Módulos Técnicos ─────────────────────────────────────────────────────

def _append_section(ws, row, titulo, linhas, bg=DARK_NAVY):
    """Insere uma seção de dados na aba de módulos técnicos."""
    _section_header(ws, row, 1, 4, titulo, bg=bg)
    row += 1

    for label, value, status, mensagem in linhas:
        ws.row_dimensions[row].height = 16
        valores = [label, _excel_value(value), _excel_value(status), _excel_value(mensagem)]
        row_fill = SURFACE
        for col, cell_value in enumerate(valores, 1):
            cell = ws.cell(row=row, column=col, value=cell_value)
            if col == 1:
                _body(cell, fill=LIGHT_BLUE, bold=True, align="left", color=DARK_NAVY)
            elif col == 3 and status:
                _status_cell(cell, status, bold=True)
            elif col == 4:
                _body(cell, fill=row_fill, align="left", color=SLATE)
            else:
                _body(cell, fill=row_fill, align="left", color=DARK_NAVY)
        row += 1
    return row + 1


def _aba_modulos_tecnicos(wb, projeto, circuitos):
    ws = wb.create_sheet("Módulos Técnicos")
    ws.sheet_properties.tabColor = PETROBRAS_BLUE
    _titulo_aba(ws, "Módulos Técnicos — Dados Normativos", "Transformador, sistema elétrico, proteções, para-raios, aterramento e áreas classificadas.", 4)

    # Sub-cabeçalho de colunas
    ws.row_dimensions[4].height = 22
    for col, header in enumerate(["Item", "Valor", "Status", "Observação"], 1):
        _header(ws.cell(row=4, column=col, value=header))

    ws.freeze_panes = "A5"

    transformador = calcular_transformador(carregar_json(getattr(projeto, "transformador_dados", None)))
    sistema = calcular_sistema_trifasico_projeto(
        carregar_json(getattr(projeto, "sistema_trifasico_dados", None)),
        circuitos,
        getattr(projeto, "tensao_ref", None),
    )
    protecao = validar_protecao_geral(carregar_json(getattr(projeto, "protecao_geral_dados", None)), transformador)
    para_raios = calcular_para_raios(carregar_json(getattr(projeto, "para_raios_dados", None)))
    aterramento = calcular_aterramento(carregar_json(getattr(projeto, "aterramento_dados", None)))
    areas = calcular_areas_classificadas(carregar_json(getattr(projeto, "areas_classificadas_dados", None)))
    resumo_global = _resumo_global(projeto, circuitos, {
        "Proteção geral": protecao,
        "Para-raios": para_raios,
        "Aterramento": aterramento,
        "Áreas Classificadas": areas,
    })
    protecoes = resumo_global["protecoes"]

    row = 6
    row = _append_section(ws, row, "Status geral do projeto", [
        ("Status geral", _display_status(resumo_global["status"]), resumo_global["status"], resumo_global["motivo_status"]),
        ("Módulos OK", resumo_global["modulos_contagem"]["ok"], "OK", ""),
        ("Módulos em ALERTA", resumo_global["modulos_contagem"]["alerta"], "ALERTA" if resumo_global["modulos_contagem"]["alerta"] else None, ""),
        ("Módulos CRÍTICO", resumo_global["modulos_contagem"]["critico"], "CRITICO" if resumo_global["modulos_contagem"]["critico"] else None, ""),
    ], bg=DARK_NAVY)

    row = _append_section(ws, row, "Entrada / Transformador", [
        ("Potência nominal", transformador.get("potencia_kva"), "OK" if transformador.get("potencia_kva") else "ALERTA", "kVA"),
        ("Tensão primária", transformador.get("tensao_primaria_formatada"), "OK" if transformador.get("tensao_primaria") else "ALERTA", ""),
        ("Tensão secundária", transformador.get("tensao_secundaria_formatada"), "OK" if transformador.get("tensao_secundaria") else "ALERTA", ""),
        ("In primário", transformador.get("corrente_nominal_primario"), None, "A"),
        ("In secundário", transformador.get("corrente_nominal_secundario"), None, "A"),
        ("Icc secundário", transformador.get("corrente_curto_secundario_ka"), None, "kA"),
    ], bg=SUBHEADER_ROW)

    row = _append_section(ws, row, "Sistema Elétrico", [
        ("Potência nominal das cargas", sistema.get("potencia_nominal_cargas_kw"), "OK" if sistema.get("potencia_nominal_cargas_kw") else "ALERTA", "kW"),
        ("Potência ativa elétrica estimada", sistema.get("potencia_ativa_eletrica_kw") or sistema.get("potencia_ativa_kw"), "OK" if sistema.get("potencia_ativa_eletrica_kw") or sistema.get("potencia_ativa_kw") else "ALERTA", "kW"),
        ("Potência aparente", sistema.get("potencia_aparente_kva"), "OK" if sistema.get("potencia_aparente_kva") else "ALERTA", "kVA"),
        ("Corrente de linha", sistema.get("corrente_linha"), "OK" if sistema.get("corrente_linha") else "ALERTA", "A"),
        ("Fator de potência", sistema.get("fator_potencia"), None, ""),
    ], bg=SUBHEADER_ROW)

    row = _append_section(ws, row, "Proteção Geral", [
        ("Disjuntor geral TAG", protecao.get("tag"), protecao.get("status"), ""),
        ("Disjuntor geral tipo", protecao.get("tipo"), protecao.get("status"), ""),
        ("Disjuntor geral Vn", protecao.get("vn"), protecao.get("status"), "V"),
        ("Disjuntor geral In", protecao.get("in"), protecao.get("status"), "A"),
        ("Disjuntor geral Icu", protecao.get("icu"), protecao.get("status"), "kA"),
        ("Disjuntor geral curva", protecao.get("curva"), protecao.get("status"), ""),
        ("Status disjuntor geral", _display_status(protecao.get("status")), protecao.get("status"), protecao.get("mensagem")),
        ("Proteções de circuitos OK", protecoes["ok"], "OK" if protecoes["ok"] else None, ""),
        ("Proteções de circuitos ALERTA", protecoes["alerta"], "ALERTA" if protecoes["alerta"] else None, ""),
        ("Proteções de circuitos CRÍTICO", protecoes["critico"], "CRITICO" if protecoes["critico"] else None, ""),
    ], bg=SUBHEADER_ROW)

    row = _append_section(ws, row, "Para-raios", [
        ("TAG", para_raios.get("tag"), para_raios.get("status"), para_raios.get("mensagem")),
        ("Vn mínimo", para_raios.get("vn_minimo"), para_raios.get("status"), ""),
        ("Vn escolhido", para_raios.get("tensao_nominal_escolhida_vn"), para_raios.get("status"), ""),
        ("Margem de proteção", para_raios.get("margem_protecao_pct"), para_raios.get("status"), "%"),
    ], bg=SUBHEADER_ROW)

    row = _append_section(ws, row, "Aterramento", [
        ("Tipo", aterramento.get("tipo_aterramento"), aterramento.get("status"), ""),
        ("Finalidade", aterramento.get("finalidade"), aterramento.get("status"), ""),
        ("Medições Wenner", len(aterramento.get("medicoes") or []), aterramento.get("status"), aterramento.get("observacoes_tecnicas")),
        ("Resistividade média", aterramento.get("resistividade_media"), aterramento.get("status"), "ohm.m"),
        ("Classificação do solo", aterramento.get("classificacao_solo"), aterramento.get("status"), aterramento.get("mensagem")),
    ], bg=SUBHEADER_ROW)

    row = _append_section(ws, row, "Áreas Classificadas", [
        ("Total de áreas", areas.get("resumo", {}).get("total_areas"), areas.get("status"), areas.get("mensagem")),
        ("Equipamentos Ex", areas.get("resumo", {}).get("total_equipamentos"), areas.get("status"), ""),
        ("Equipamentos críticos", areas.get("resumo", {}).get("equipamentos_criticos"), areas.get("status"), ""),
    ], bg=SUBHEADER_ROW)

    for area in areas.get("areas") or []:
        row = _append_section(ws, row, f"Área Ex — {area.get('nome') or 'sem identificação'}", [
            ("Classificação", f"{area.get('zona') or 'N/D'} / {area.get('grupo') or 'N/D'} / {area.get('classe_temperatura') or 'N/D'}", area.get("status"), area.get("mensagem")),
            ("Equipamentos cadastrados", len(area.get("equipamentos") or []), area.get("status"), area.get("observacoes")),
        ], bg="475569")
        for equipamento in area.get("equipamentos") or []:
            row = _append_section(ws, row, f"Equipamento Ex — {equipamento.get('nome_tag') or 'sem tag'}", [
                ("Tipo", equipamento.get("tipo"), equipamento.get("status"), ""),
                ("Proteção Ex", equipamento.get("tipo_protecao_ex") or "nenhuma", equipamento.get("status"), ""),
                ("Grupo / Temperatura", f"{equipamento.get('grupo') or 'N/D'} / {equipamento.get('classe_temperatura') or 'N/D'}", equipamento.get("status"), ""),
                ("Certificado", equipamento.get("certificado") or "não informado", equipamento.get("status"), equipamento.get("mensagem")),
            ], bg="64748B")

    if resumo_global["status"] == "CRITICO":
        conclusao = "Projeto NÃO está liberado para emissão final pelo CalcCabos."
    elif resumo_global["status"] == "ALERTA":
        conclusao = "Projeto requer revisão técnica antes da emissão final."
    else:
        conclusao = "Projeto atende aos critérios de dimensionamento calculados pelo CalcCabos."

    pendencias = []
    if normalizar_status(para_raios.get("status"), "ALERTA") == "ALERTA":
        pendencias.append("Para-raios não informado ou incompleto.")
    if normalizar_status(aterramento.get("status"), "ALERTA") == "ALERTA":
        pendencias.append(aterramento.get("mensagem") or "Aterramento em alerta.")
    if normalizar_status(areas.get("status"), "ALERTA") == "CRITICO":
        pendencias.append(areas.get("mensagem") or "Áreas Classificadas com criticidade.")

    row = _append_section(ws, row, "Conclusão técnica", [
        ("Conclusão", conclusao, resumo_global["status"], "; ".join(pendencias) if pendencias else "Sem pendências críticas identificadas."),
    ], bg=DARK_NAVY)

    _ajustar_larguras(ws, [32, 28, 18, 62])
    return ws


# ─── Aba Pendências ───────────────────────────────────────────────────────────

def _aba_pendencias(wb, projeto, circuitos):
    """Nova aba consolidando todas as pendências do projeto."""
    ws = wb.create_sheet("Pendências")
    ws.sheet_properties.tabColor = CRITICAL_FILL

    resumo_global = _resumo_global(projeto, circuitos)
    transformador = calcular_transformador(carregar_json(getattr(projeto, "transformador_dados", None)))
    protecao = validar_protecao_geral(carregar_json(getattr(projeto, "protecao_geral_dados", None)), transformador)
    para_raios = calcular_para_raios(carregar_json(getattr(projeto, "para_raios_dados", None)))
    aterramento = calcular_aterramento(carregar_json(getattr(projeto, "aterramento_dados", None)))
    areas = calcular_areas_classificadas(carregar_json(getattr(projeto, "areas_classificadas_dados", None)))

    _titulo_aba(ws, "Pendências e Não-Conformidades", "Itens que requerem revisão técnica antes da emissão final.", 6)

    header_row = 4
    ws.row_dimensions[header_row].height = 26
    headers = ["Tipo", "Módulo / TAG", "Circuito / Item", "Status", "Mensagem técnica", "Recomendação"]
    for col, h in enumerate(headers, 1):
        _header(ws.cell(row=header_row, column=col, value=h))

    row = 5
    pendencias = []

    # Circuitos com status CRÍTICO ou ALERTA
    for c in circuitos:
        st = _status_final(c)
        if st in ("CRITICO", "ALERTA"):
            tag = c.tag or c.descricao or "-"
            msg = getattr(c, "validacao_mensagem", None) or getattr(c, "selecao_componentes_justificativa", None) or "-"
            rec = "Revisar dimensionamento, queda de tensão e proteção do circuito."
            pendencias.append(("Circuito", tag, tag, st, msg, rec))

    # Proteção de circuitos
    for c in circuitos:
        pst = normalizar_status(getattr(c, "protecao_status", None), "OK")
        if pst in ("CRITICO", "ALERTA"):
            tag = c.tag or c.descricao or "-"
            nota = getattr(c, "protecao_nota", None) or "-"
            pendencias.append(("Proteção", tag, tag, pst, nota, "Verificar In, Icu e curva do disjuntor do circuito."))

    # Módulos
    modulos_check = [
        ("Proteção Geral", protecao),
        ("Para-raios", para_raios),
        ("Aterramento", aterramento),
        ("Áreas Classificadas", areas),
    ]
    for nome, dados in modulos_check:
        st = normalizar_status(dados.get("status"), "ALERTA")
        if st in ("CRITICO", "ALERTA"):
            msg = dados.get("mensagem") or "-"
            pendencias.append(("Módulo", nome, nome, st, msg, f"Revisar dados do módulo {nome}."))

    if not pendencias:
        ws.merge_cells("A5:F5")
        ws["A5"] = "Nenhuma pendência identificada. Projeto aprovado."
        ws["A5"].font = Font(bold=True, color=OK_TEXT, size=10, name="Calibri")
        ws["A5"].fill = PatternFill("solid", fgColor=OK_FILL)
        ws["A5"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[5].height = 28
    else:
        for idx, (tipo, modulo, circuito, status, msg, rec) in enumerate(pendencias):
            ws.row_dimensions[row].height = 16
            row_fill = LIGHT_GRAY if idx % 2 == 0 else WHITE
            values = [tipo, modulo, circuito, _display_status(status), msg, rec]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=value)
                if col == 4:
                    _status_cell(cell, status, bold=True)
                elif col in (1, 2, 3):
                    _body(cell, fill=row_fill, bold=(col == 1), align="left", color=DARK_NAVY)
                else:
                    _body(cell, fill=row_fill, align="left", color=SLATE)
            row += 1

    ws.auto_filter.ref = f"A{header_row}:F{header_row}"
    ws.freeze_panes = "A5"
    _ajustar_larguras(ws, [16, 22, 22, 14, 55, 48])
    return ws


# ─── Aba Memorial Cabos (Nova Aba Simplificada) ───────────────────────────────

def _aba_memorial_cabos(wb, projeto, circuitos):
    ws = wb.create_sheet("Memorial Cabos")
    ws.sheet_properties.tabColor = ACCENT_BLUE
    _titulo_aba(ws, "Memorial Cabos", "Memorial simplificado com os dados consolidados do projeto.", 26)

    memorial = montar_memorial_cabos(projeto, circuitos)

    header_row = 4
    ws.row_dimensions[header_row].height = 26
    
    headers = [
        "TAG", "Descrição", "De", "Para", "V", "Fases", "kW", "kVA", "FP", "η",
        "Ib", "Ib Corrigida", "K1", "K2", "K3", "Método", "Temperatura", "Agrupamento",
        "Cabo", "Ampacidade", "ΔV", "STATUS_GERAL", "ICC_LOCAL_KA", "DISJUNTOR", "CURVA", "ICU_KA", "STATUS_PROTECAO", "JUSTIFICATIVA_PROTECAO"
    ]
    
    for col, header in enumerate(headers, 1):
        _header(ws.cell(row=header_row, column=col, value=header))

    row = 5
    for c_dados in memorial["circuitos"]:
        ws.row_dimensions[row].height = 16
        row_fill = LIGHT_GRAY if (row - 4) % 2 == 0 else WHITE
        
        values = [
            c_dados["tag"],
            c_dados["descricao"],
            c_dados["origem"],
            c_dados["destino"],
            c_dados["tensao"],
            c_dados["fases"],
            c_dados["potencia_kw"],
            c_dados["potencia_kva"],
            c_dados["fator_potencia"],
            c_dados["eficiencia"],
            c_dados["corrente_projeto_ib"],
            c_dados["corrente_corrigida_ib_linha"],
            c_dados["fator_k1"],
            c_dados["fator_k2"],
            c_dados["fator_k3"],
            c_dados["metodo_instalacao"],
            c_dados["temperatura_ambiente"],
            c_dados["agrupamento"],
            c_dados["cabo_selecionado"],
            c_dados["ampacidade_icond"],
            c_dados["queda_tensao"],
            c_dados["status_final"],
            c_dados["corrente_curto_icc"],
            c_dados["disjuntor_in"],
            c_dados["disjuntor_curva"],
            c_dados["disjuntor_icu"],
            c_dados["status_protecao"],
            c_dados["justificativa_protecao"]
        ]
        
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=_excel_value(value))
            if col in [22, 27]:  # STATUS_GERAL, STATUS_PROTECAO
                _status_cell(cell, value, bold=True)
            elif col in [1, 2, 28]:  # TAG, Descrição, Justificativa
                _body(cell, fill=row_fill, align="left")
            else:
                _body(cell, fill=row_fill)
        row += 1

    ws.auto_filter.ref = f"A{header_row}:AB{header_row}"
    ws.freeze_panes = "A5"
    _ajustar_larguras(ws, [
        12, 25, 15, 15, 8, 8, 8, 8, 8, 8,
        10, 12, 8, 8, 8, 12, 12, 12,
        16, 12, 8, 14, 14, 12, 10, 10, 16, 40
    ])
    return ws


# ─── Função principal ─────────────────────────────────────────────────────────

def gerar_excel(projeto, circuitos, usuario):
    wb = Workbook()
    try:
        wb.calculation.fullCalcOnLoad = True
    except AttributeError:
        pass

    _aba_memoria(wb, projeto, circuitos, usuario)
    _aba_memorial_cabos(wb, projeto, circuitos)
    _aba_lista_cargas(wb, circuitos)
    _aba_modulos_tecnicos(wb, projeto, circuitos)
    _aba_resumo(wb, projeto, circuitos)
    _aba_pendencias(wb, projeto, circuitos)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
