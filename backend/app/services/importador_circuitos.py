import csv
import io
import re
import unicodedata
from typing import Any, Dict, List, Tuple

from openpyxl import load_workbook
import pandas as pd

from app.services.status_utils import normalizar_tipo_cabo


COLUNAS_ESPERADAS = {
    "nome": ["nome", "descricao", "descrição", "circuito", "load", "carga", "id circuito", "identificacao", "equipamento"],
    "potencia": ["potencia", "potência", "potencia kw", "potência kw", "kw", "p", "p kw", "potencia instalada", "carga kw", "kw total", "pw"],
    "tensao": ["tensao", "tensão", "tensao v", "tensão v", "v", "vn", "tensao nominal", "v nominal", "un", "tensao linha"],
    "fp": ["fp", "fator potencia", "fator de potencia", "fator de potência", "cos phi", "cosfi", "cos fi", "power factor", "pf", "fator de pot"],
    "comprimento": ["comprimento", "distancia", "distância", "distancia m", "distância m", "l", "l(m)", "l m", "comp m", "comprimento m", "distancia cabo", "metros"],
    "metodo_instalacao": ["metodo", "método", "metodo instalacao", "método instalação", "instalacao", "instalação", "instalacao cabo", "modo cabo", "routing"],
    "temperatura": ["temperatura", "temp", "temp ambiente", "temperatura ambiente"],
    "tipo_carga": ["tipo carga", "tipo de carga", "carga tipo", "load type", "categoria", "classe carga", "load class"],
    "tag": ["tag", "identificacao", "identificação"],
    "tipo_cabo": ["tipo cabo", "cabo", "isolacao", "isolação"],
    "fases": ["fases", "fase", "phase", "phases", "n fases", "num fases", "monofasico", "trifasico"],
    "agrupamento": ["agrupamento", "cabos agrupados", "grupo"],
    "formacao": ["formacao", "formação", "paralelos", "cabos paralelos"],
    "icc": ["icc", "isc", "curto", "curto circuito", "ka"],
    "tempo_atuacao": ["tempo atuacao", "tempo atuação", "t(s)", "tempo", "atuacao"],
    "disjuntor_corrente_nominal": ["in disjuntor", "disjuntor in", "in dj", "breaker in"],
    "disjuntor_icu": ["icu", "icu disjuntor", "capacidade interrupcao", "capacidade de interrupcao"],
    "disjuntor_curva": ["curva", "curva disjuntor", "trip curve"],
    "disjuntor_fabricante": ["fabricante", "fabricante disjuntor", "marca"],
}

CAMPOS_OBRIGATORIOS = ["nome", "potencia", "tensao", "fp", "comprimento"]


def _normalizar_texto(valor: Any) -> str:
    texto = str(valor or "").strip().lower()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(ch for ch in texto if not unicodedata.combining(ch))
    texto = re.sub(r"[^a-z0-9]+", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


def _parse_float(valor: Any):
    if valor is None or valor == "":
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip()
    if not texto:
        return None
    texto = re.sub(r"[^0-9,.\-]", "", texto)
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def _parse_int(valor: Any, padrao: int):
    numero = _parse_float(valor)
    if numero is None:
        return padrao
    return int(numero)


def _normalizar_metodo(valor: Any) -> str:
    texto = _normalizar_texto(valor)
    if "eletroduto" in texto or "conduit" in texto:
        return "CONDUIT"
    if "enterr" in texto or "direct" in texto:
        return "DIRECT"
    if "ar livre" in texto or texto == "ar" or "air" in texto:
        return "AIR"
    return "TRAY"


def _normalizar_tipo_cabo(valor: Any) -> str:
    return normalizar_tipo_cabo(valor)


def lcs_length(a: str, b: str) -> int:
    m, n = len(a), len(b)
    if m == 0 or n == 0: return 0
    dp = [[0]*(n+1) for _ in range(m+1)]
    for i in range(1, m+1):
        for j in range(1, n+1):
            if a[i-1] == b[j-1]:
                dp[i][j] = dp[i-1][j-1] + 1
            else:
                dp[i][j] = max(dp[i-1][j], dp[i][j-1])
    return dp[m][n]

def _detectar_colunas(headers: List[Any]) -> Tuple[Dict[str, str], Dict[str, str]]:
    headers_norm = {_normalizar_texto(h): str(h).strip() for h in headers if str(h or "").strip()}
    detectadas = {}
    confianca = {}
    ordem_campos = list(COLUNAS_ESPERADAS.keys())
    for header_norm, header_original in headers_norm.items():
        best_campo = None
        best_camada = 5
        best_indice = 999
        for campo, aliases in COLUNAS_ESPERADAS.items():
            indice_campo = ordem_campos.index(campo)
            aliases_norm = [_normalizar_texto(alias) for alias in aliases]
            camada = 5
            for alias in aliases_norm:
                if header_norm == alias:
                    camada = min(camada, 1)
                tokens_alias = alias.split()
                tokens_header = header_norm.split()
                if tokens_alias and all(t in tokens_header for t in tokens_alias):
                    camada = min(camada, 2)
                if len(alias) >= 4:
                    score = lcs_length(header_norm, alias) / max(len(header_norm), len(alias))
                    if score >= 0.75:
                        camada = min(camada, 3)
                if len(min(header_norm, alias, key=len)) >= 3:
                    if alias in header_norm or header_norm in alias:
                        camada = min(camada, 4)
            if camada < 5:
                if camada < best_camada or (camada == best_camada and indice_campo < best_indice):
                    best_camada = camada
                    best_campo = campo
                    best_indice = indice_campo
        if best_campo and best_campo not in detectadas:
            detectadas[best_campo] = header_original
            if best_camada <= 2:
                confianca[best_campo] = "alto"
            elif best_camada == 3:
                confianca[best_campo] = "medio"
            else:
                confianca[best_campo] = "baixo"
    return detectadas, confianca


def _ler_xlsx(conteudo: bytes) -> Tuple[List[str], List[Dict[str, Any]], str]:
    wb = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    melhor_aba = wb.sheetnames[0]
    best_score = -1
    palavras_chave = ["circuito", "circuit", "carga", "load", "dados", "data"]
    for sheet_name in wb.sheetnames:
        nome_norm = _normalizar_texto(sheet_name)
        if any(kw in nome_norm for kw in palavras_chave):
            melhor_aba = sheet_name
            best_score = 999
            break
    if best_score < 999:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            linhas = list(ws.iter_rows(values_only=True, max_row=10))
            if not linhas: continue
            for linha in linhas:
                str_linha = [str(v) for v in linha if v is not None and str(v).strip()]
                if not str_linha: continue
                det, _ = _detectar_colunas(str_linha)
                if len(det) > best_score:
                    best_score = len(det)
                    melhor_aba = sheet_name
    ws = wb[melhor_aba]
    linhas = list(ws.iter_rows(values_only=True))
    headers, registros = _linhas_para_dicts(linhas)
    return headers, registros, melhor_aba

def _ler_xls(conteudo: bytes) -> Tuple[List[str], List[Dict[str, Any]], str]:
    df_dict = pd.read_excel(io.BytesIO(conteudo), sheet_name=None, engine="xlrd", header=None)
    sheet_names = list(df_dict.keys())
    melhor_aba = sheet_names[0]
    best_score = -1
    palavras_chave = ["circuito", "circuit", "carga", "load", "dados", "data"]
    for sheet_name in sheet_names:
        nome_norm = _normalizar_texto(sheet_name)
        if any(kw in nome_norm for kw in palavras_chave):
            melhor_aba = sheet_name
            best_score = 999
            break
    if best_score < 999:
        for sheet_name in sheet_names:
            df = df_dict[sheet_name]
            linhas = df.head(10).values.tolist()
            for linha in linhas:
                str_linha = [str(v) for v in linha if pd.notna(v) and str(v).strip()]
                if not str_linha: continue
                det, _ = _detectar_colunas(str_linha)
                if len(det) > best_score:
                    best_score = len(det)
                    melhor_aba = sheet_name
    df = df_dict[melhor_aba]
    linhas = df.where(pd.notna(df), None).values.tolist()
    headers, registros = _linhas_para_dicts(linhas)
    return headers, registros, melhor_aba

def _ler_csv(conteudo: bytes) -> Tuple[List[str], List[Dict[str, Any]], str]:
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("latin-1")
    amostra = texto[:2048]
    try:
        dialecto = csv.Sniffer().sniff(amostra, delimiters=";,")
        delimitador = dialecto.delimiter
    except csv.Error:
        delimitador = ";" if texto.count(";") >= texto.count(",") else ","
    linhas = list(csv.reader(io.StringIO(texto), delimiter=delimitador))
    headers, registros = _linhas_para_dicts(linhas)
    return headers, registros, "csv_aba"

def _linhas_para_dicts(linhas: List[Any]) -> Tuple[List[str], List[Dict[str, Any]]]:
    header_index = None
    headers = []
    melhor_index = None
    melhor_count = -1
    melhor_headers = []
    for idx, linha in enumerate(linhas[:10]):
        valores = [str(v).strip() for v in linha if v is not None and str(v).strip()]
        if not valores: continue
        det, _ = _detectar_colunas(valores)
        if len(det) >= 2 and len(det) > melhor_count:
            melhor_count = len(det)
            melhor_index = idx
            melhor_headers = [str(v).strip() if v is not None else "" for v in linha]
    if melhor_index is not None:
        header_index = melhor_index
        headers = melhor_headers
    else:
        for idx, linha in enumerate(linhas):
            valores = [str(v).strip() for v in linha if v is not None and str(v).strip()]
            if len(valores) >= 2:
                header_index = idx
                headers = [str(v).strip() if v is not None else "" for v in linha]
                break
    if header_index is None:
        raise ValueError("Arquivo sem cabecalho detectavel.")

    registros = []
    for numero, linha in enumerate(linhas[header_index + 1 :], start=header_index + 2):
        if not any(v is not None and str(v).strip() for v in linha):
            continue
        registro = {headers[i]: linha[i] if i < len(linha) else None for i in range(len(headers)) if headers[i]}
        registro["_linha"] = numero
        registros.append(registro)
    return headers, registros


def ler_arquivo(nome_arquivo: str, conteudo: bytes) -> Tuple[List[str], List[Dict[str, Any]], str]:
    nome = nome_arquivo.lower()
    if nome.endswith(".xlsx"):
        return _ler_xlsx(conteudo)
    if nome.endswith(".xls"):
        return _ler_xls(conteudo)
    if nome.endswith(".csv"):
        return _ler_csv(conteudo)
    raise ValueError("Formato nao suportado. Envie arquivo .xlsx, .xls ou .csv.")


def _valor(registro: Dict[str, Any], mapa: Dict[str, str], campo: str):
    coluna = mapa.get(campo)
    if not coluna:
        return None
    return registro.get(coluna)


def _validar_registro(registro: Dict[str, Any], mapa: Dict[str, str]) -> Dict[str, Any]:
    erros = []
    linha = registro.get("_linha")

    for campo in CAMPOS_OBRIGATORIOS:
        if campo not in mapa:
            erros.append(f"Coluna obrigatoria ausente: {campo}")

    nome = str(_valor(registro, mapa, "nome") or "").strip()
    potencia = _parse_float(_valor(registro, mapa, "potencia"))
    tensao = _parse_float(_valor(registro, mapa, "tensao"))
    fp = _parse_float(_valor(registro, mapa, "fp"))
    comprimento = _parse_float(_valor(registro, mapa, "comprimento"))
    temperatura = _parse_float(_valor(registro, mapa, "temperatura")) or 30.0
    tipo_carga = str(_valor(registro, mapa, "tipo_carga") or "").strip()

    if not nome:
        erros.append("nome obrigatorio")
    if potencia is None or potencia <= 0:
        erros.append("potencia deve ser maior que zero")
    if tensao is None or tensao <= 0:
        erros.append("tensao deve ser maior que zero")
    if fp is None or fp <= 0 or fp > 1:
        erros.append("fp deve estar entre 0 e 1")
    if comprimento is None or comprimento <= 0:
        erros.append("comprimento deve ser maior que zero")

    metodo = _normalizar_metodo(_valor(registro, mapa, "metodo_instalacao"))
    tipo_cabo = _normalizar_tipo_cabo(_valor(registro, mapa, "tipo_cabo"))
    tipo_carga_norm = _normalizar_texto(tipo_carga)
    eficiencia = 0.92 if "motor" in tipo_carga_norm else 1.0
    protection_device = "MCCB" if "motor" in tipo_carga_norm else "DJ"
    nota = f"Tipo de carga importado: {tipo_carga}" if tipo_carga else None

    circuito = None
    if not erros:
        circuito = {
            "descricao": nome,
            "tag": str(_valor(registro, mapa, "tag") or "").strip() or None,
            "tensao": tensao,
            "potencia_kw": potencia,
            "fator_potencia": fp,
            "distancia_m": comprimento,
            "comprimento_real": comprimento,
            "metodo_instalacao": metodo,
            "temp_ambiente": temperatura,
            "tipo_cabo": tipo_cabo,
            "fases": _parse_int(_valor(registro, mapa, "fases"), 3),
            "agrupamento": max(_parse_int(_valor(registro, mapa, "agrupamento"), 1), 1),
            "formacao": max(_parse_int(_valor(registro, mapa, "formacao"), 1), 1),
            "fator_eficiencia": eficiencia,
            "fator_demanda": 1.0,
            "corrente_ac_dc": "AC",
            "protection_device": protection_device,
            "isc_local": _parse_float(_valor(registro, mapa, "icc")),
            "tempo_atuacao": _parse_float(_valor(registro, mapa, "tempo_atuacao")) or 0.1,
            "disjuntor_corrente_nominal": _parse_float(_valor(registro, mapa, "disjuntor_corrente_nominal")),
            "disjuntor_icu": _parse_float(_valor(registro, mapa, "disjuntor_icu")),
            "disjuntor_curva": str(_valor(registro, mapa, "disjuntor_curva") or "").strip().upper() or None,
            "disjuntor_fabricante": str(_valor(registro, mapa, "disjuntor_fabricante") or "").strip() or None,
            "nota_tecnica": nota,
        }

    return {
        "linha": linha,
        "valido": not erros,
        "erros": erros,
        "dados_originais": {k: v for k, v in registro.items() if k != "_linha"},
        "circuito": circuito,
    }


# Power BI Import Flow — score por caracteres em comum entre dois textos normalizados
def _score_similaridade(a: str, b: str) -> int:
    """Conta quantos caracteres (com repetição) a e b têm em comum."""
    from collections import Counter
    ca, cb = Counter(a), Counter(b)
    return sum((ca & cb).values())


def _sugestoes_colunas(headers: List[str]) -> Dict[str, str]:
    """
    Para cada coluna do arquivo, retorna o campo do sistema mais próximo
    com base na quantidade de caracteres em comum (score).
    Retorna dict: {coluna_arquivo: campo_sistema}
    """
    sugestoes: Dict[str, str] = {}
    todos_aliases: Dict[str, List[str]] = {
        campo: [_normalizar_texto(a) for a in aliases]
        for campo, aliases in COLUNAS_ESPERADAS.items()
    }

    for header in headers:
        if not str(header or "").strip():
            continue
        header_norm = _normalizar_texto(header)
        melhor_campo = None
        melhor_score = 0
        for campo, aliases_norm in todos_aliases.items():
            for alias in aliases_norm:
                s = _score_similaridade(header_norm, alias)
                if s > melhor_score:
                    melhor_score = s
                    melhor_campo = campo
        if melhor_campo and melhor_score >= 2:
            sugestoes[str(header).strip()] = melhor_campo
    return sugestoes


def preview_importacao(nome_arquivo: str, conteudo: bytes) -> Dict[str, Any]:
    headers, registros, aba_usada = ler_arquivo(nome_arquivo, conteudo)
    mapa, confianca = _detectar_colunas(headers)
    linhas = [_validar_registro(registro, mapa) for registro in registros]
    validas = sum(1 for linha in linhas if linha["valido"])
    invalidas = len(linhas) - validas

    # Power BI Import Flow — campos extras para o wizard de mapeamento
    headers_limpos = [h for h in headers if h]

    amostra_raw = [
        {k: v for k, v in r.items() if k != "_linha"}
        for r in registros[:5]
    ]

    colunas_nao_mapeadas = [
        campo for campo in CAMPOS_OBRIGATORIOS if campo not in mapa
    ]
    
    colunas_nao_mapeadas_arquivo = [
        h for h in headers_limpos if h not in mapa.values()
    ]

    sugestoes = _sugestoes_colunas(headers_limpos)
    total_linhas_arquivo = registros[-1]["_linha"] if registros else 0

    return {
        "arquivo": nome_arquivo,
        "colunas": headers,
        "mapeamento": mapa,
        "linhas": linhas,
        "resumo": {"total": len(linhas), "validas": validas, "invalidas": invalidas},
        # Novos campos Power BI Import Flow
        "amostra": amostra_raw,
        "colunas_detectadas": mapa,
        "colunas_nao_mapeadas": colunas_nao_mapeadas,
        "sugestoes": sugestoes,
        "aba_usada": aba_usada,
        "total_linhas_arquivo": total_linhas_arquivo,
        "colunas_arquivo": headers_limpos,
        "mapeamento_automatico": mapa,
        "confianca_mapeamento": confianca,
        "colunas_nao_mapeadas_arquivo": colunas_nao_mapeadas_arquivo,
        "campos_obrigatorios_faltando": colunas_nao_mapeadas,
        "amostra_raw": amostra_raw,
    }


def confirmar_importacao(
    nome_arquivo: str,
    conteudo: bytes,
    mapeamento_usuario: Dict[str, str],
) -> Dict[str, Any]:
    """
    Power BI Import Flow — usa o mapeamento confirmado pelo usuário
    (campo_sistema → coluna_xlsx) em vez da detecção automática.
    Mantém toda a validação e estrutura de retorno existentes.
    """
    headers, registros, aba_usada = ler_arquivo(nome_arquivo, conteudo)

    # Filtra apenas entradas com coluna válida (ignora "-- não mapear --" e vazios)
    mapa = {
        campo: coluna
        for campo, coluna in mapeamento_usuario.items()
        if coluna and coluna not in ("", "--")
    }

    linhas = [_validar_registro(registro, mapa) for registro in registros]
    validas = sum(1 for linha in linhas if linha["valido"])
    invalidas = len(linhas) - validas

    colunas_nao_mapeadas = [
        campo for campo in CAMPOS_OBRIGATORIOS if campo not in mapa
    ]

    return {
        "arquivo": nome_arquivo,
        "colunas": headers,
        "mapeamento": mapa,
        "linhas": linhas,
        "resumo": {"total": len(linhas), "validas": validas, "invalidas": invalidas},
        "colunas_nao_mapeadas": colunas_nao_mapeadas,
    }
